import streamlit as st
import re
import gspread
from google.oauth2 import service_account
from datetime import date

st.set_page_config(
    page_title="新薬メンテナンス",
    page_icon="🧪",
    layout="centered",
    initial_sidebar_state="expanded",
    menu_items={}
)

st.sidebar.title("メニュー")

SPREADSHEET_URL = "https://docs.google.com/spreadsheets/d/1dLEUYSZlrIK1uHqEtEAfS1jSAPpXCIiAiAk_iaRuY-8/edit"

MASTER_HEADERS = [
    "管理コード", "一般名（全角）", "一般名（半角カナ）",
    "採用商品名（全角）", "採用商品名（半角カナ）",
    "薬効分類", "薬剤区分", "支持療法分類", "単位", "投与経路",
    "標準希釈液", "フィルター", "遮光", "先発後発", "備考",
    "1V当たりmg", "患者向け説明", "スケジュールシール用種類",
    "別名・旧採用品名", "短縮注記",
]

PREFIX_OPTIONS = {
    "AC（抗がん剤・注射）": "AC",
    "ACO（抗がん剤・内服）": "ACO",
    "SJ（支持療法・注射）": "SJ",
    "SO（支持療法・内服）": "SO",
    "IV（輸液）": "IV",
}
PREFIX_DEFAULT_KUBUN = {
    "AC": "抗がん剤",
    "ACO": "抗がん剤",
    "SJ": "支持療法",
    "SO": "支持療法",
    "IV": "輸液",
}

TO_HALF_KANA_TABLE = {
    'ア':'ｱ','イ':'ｲ','ウ':'ｳ','エ':'ｴ','オ':'ｵ',
    'カ':'ｶ','キ':'ｷ','ク':'ｸ','ケ':'ｹ','コ':'ｺ',
    'サ':'ｻ','シ':'ｼ','ス':'ｽ','セ':'ｾ','ソ':'ｿ',
    'タ':'ﾀ','チ':'ﾁ','ツ':'ﾂ','テ':'ﾃ','ト':'ﾄ',
    'ナ':'ﾅ','ニ':'ﾆ','ヌ':'ﾇ','ネ':'ﾈ','ノ':'ﾉ',
    'ハ':'ﾊ','ヒ':'ﾋ','フ':'ﾌ','ヘ':'ﾍ','ホ':'ﾎ',
    'マ':'ﾏ','ミ':'ﾐ','ム':'ﾑ','メ':'ﾒ','モ':'ﾓ',
    'ヤ':'ﾔ','ユ':'ﾕ','ヨ':'ﾖ',
    'ラ':'ﾗ','リ':'ﾘ','ル':'ﾙ','レ':'ﾚ','ロ':'ﾛ',
    'ワ':'ﾜ','ヲ':'ｦ','ン':'ﾝ',
    'ァ':'ｧ','ィ':'ｨ','ゥ':'ｩ','ェ':'ｪ','ォ':'ｫ',
    'ッ':'ｯ','ャ':'ｬ','ュ':'ｭ','ョ':'ｮ',
    'ガ':'ｶﾞ','ギ':'ｷﾞ','グ':'ｸﾞ','ゲ':'ｹﾞ','ゴ':'ｺﾞ',
    'ザ':'ｻﾞ','ジ':'ｼﾞ','ズ':'ｽﾞ','ゼ':'ｾﾞ','ゾ':'ｿﾞ',
    'ダ':'ﾀﾞ','ヂ':'ﾁﾞ','ヅ':'ﾂﾞ','デ':'ﾃﾞ','ド':'ﾄﾞ',
    'バ':'ﾊﾞ','ビ':'ﾋﾞ','ブ':'ﾌﾞ','ベ':'ﾍﾞ','ボ':'ﾎﾞ',
    'パ':'ﾊﾟ','ピ':'ﾋﾟ','プ':'ﾌﾟ','ペ':'ﾍﾟ','ポ':'ﾎﾟ',
    'ー':'ｰ','ヴ':'ｳﾞ','・':'･',
}

def preview_half_kana(text):
    """表示用プレビューのみ。実際の保存はASC()数式で行う。"""
    return ''.join(TO_HALF_KANA_TABLE.get(c, c) for c in str(text))


@st.cache_resource
def get_gspread_client():
    creds = service_account.Credentials.from_service_account_info(
        st.secrets["gcp_service_account"],
        scopes=[
            "https://spreadsheets.google.com/feeds",
            "https://www.googleapis.com/auth/drive",
        ]
    )
    return gspread.authorize(creds)

@st.cache_data(ttl=60)
def fetch_sheet_realtime(sheet_name):
    try:
        gc = get_gspread_client()
        ss = gc.open_by_url(SPREADSHEET_URL)
        ws = ss.worksheet(sheet_name)
        return ws.get_all_records()
    except Exception as e:
        st.error(f"シート「{sheet_name}」の取得に失敗: {e}")
        return []

@st.cache_data(ttl=60)
def load_all_data():
    drug_data   = fetch_sheet_realtime("薬剤情報")
    master_data = fetch_sheet_realtime("薬品マスタ")
    ae_data     = fetch_sheet_realtime("抗がん剤副作用マスタ")
    pd_data     = fetch_sheet_realtime("Pd")
    return drug_data, master_data, ae_data, pd_data


def get_next_code(prefix, master_data):
    """既存コードの表記（ハイフン有無）を検出し、次の連番コードを生成する"""
    max_num = 0
    use_hyphen = False
    pattern = re.compile(rf'^{re.escape(prefix)}-?(\d+)$')
    for m in master_data:
        code = str(m.get('管理コード', '')).strip()
        match = pattern.match(code)
        if match:
            num = int(match.group(1))
            if num > max_num:
                max_num = num
            if '-' in code:
                use_hyphen = True
    next_num = max_num + 1
    if use_hyphen:
        return f"{prefix}-{next_num:03d}"
    return f"{prefix}{next_num:03d}"


def get_unique_values(master_data, column_name):
    """薬品マスタから指定列のユニーク値一覧を取得（空欄・重複除く）"""
    values = []
    seen = set()
    for m in master_data:
        v = str(m.get(column_name, '')).strip()
        if v and v not in seen:
            seen.add(v)
            values.append(v)
    return sorted(values)


def get_ae_columns(ae_data):
    if not ae_data:
        return []
    excluded = {"管理コード", "一般名（全角）", "出典", "登録日"}
    return [k for k in ae_data[0].keys() if k not in excluded]


def get_unregistered_ae_drugs(master_data, ae_data):
    """薬品マスタのAC系コードのうち、副作用マスタに登録日が無いものを抽出"""
    ae_dict = {str(r.get('管理コード', '')).strip(): r for r in ae_data}
    result = []
    for m in master_data:
        code = str(m.get('管理コード', '')).strip()
        if not code.upper().startswith('AC'):
            continue
        ae_row = ae_dict.get(code, {})
        reg_date = str(ae_row.get('登録日', '')).strip()
        if not reg_date:
            result.append({
                'code': code,
                'name': str(m.get('一般名（全角）', code)).strip(),
            })
    return result


def show_ae_check_ui(code, name, ae_data, ae_columns):
    """新規・既存を問わず、1薬剤分の副作用チェックUIを表示する"""
    st.divider()
    st.subheader(f"💊 抗がん剤副作用マスタ：{name}（{code}）")
    st.info(
        f"📄 [PMDAで添付文書を検索](https://www.pmda.go.jp/PmdaSearch/iyakuSearch/) し、"
        f"「{name}」の添付文書・インタビューフォームの「重大な副作用」「その他の副作用」"
        "「モニタリング項目」等を目視確認して、該当する副作用にチェックを入れてください。"
    )

    ae_dict = {str(r.get('管理コード', '')).strip(): r for r in ae_data}
    current = ae_dict.get(code, {})

    cols = st.columns(3)
    checked = {}
    for i, col_name in enumerate(ae_columns):
        cb_key = f"maint_cb_{code}_{col_name}"
        if cb_key not in st.session_state:
            st.session_state[cb_key] = (str(current.get(col_name, '')).strip() == '○')
        with cols[i % 3]:
            checked[col_name] = st.checkbox(col_name, key=cb_key)

    source_text = st.text_input(
        "出典（例：添付文書(PMDA)、インタビューフォーム 等）",
        value=str(current.get('出典', '')).strip() or "添付文書(PMDA)",
        key=f"maint_src_{code}",
    )

    col_submit, col_skip = st.columns(2)
    with col_submit:
        if st.button(
            f"✅ {name} の副作用を登録する",
            type="primary",
            use_container_width=True,
            key=f"maint_ae_submit_{code}",
        ):
            try:
                gc = get_gspread_client()
                sh = gc.open_by_url(SPREADSHEET_URL)
                ws_ae = sh.worksheet("抗がん剤副作用マスタ")
                ae_all = ws_ae.get_all_values()
                ae_codes = [row[0] for row in ae_all]
                today = date.today().strftime("%Y/%m/%d")

                from openpyxl.utils import get_column_letter as gcl
                headers_ae = ae_all[0] if ae_all else []

                if code in ae_codes:
                    row_idx = ae_codes.index(code) + 1
                    # 副作用列ごとに正確な列位置へ個別書き込み（列ズレ防止）
                    for col_name in ae_columns:
                        if col_name in headers_ae:
                            col_idx = headers_ae.index(col_name) + 1
                            ws_ae.update(
                                range_name=f'{gcl(col_idx)}{row_idx}',
                                values=[['○' if checked.get(col_name, False) else '']],
                            )
                    if "出典" in headers_ae:
                        src_col = headers_ae.index("出典") + 1
                        ws_ae.update(
                            range_name=f'{gcl(src_col)}{row_idx}',
                            values=[[source_text]],
                        )
                    if "登録日" in headers_ae:
                        date_col = headers_ae.index("登録日") + 1
                        ws_ae.update(
                            range_name=f'{gcl(date_col)}{row_idx}',
                            values=[[today]],
                        )
                else:
                    # 新規行はヘッダー順に正確にマッピングして作成
                    new_row = [''] * len(headers_ae) if headers_ae else []
                    if headers_ae:
                        if "管理コード" in headers_ae:
                            new_row[headers_ae.index("管理コード")] = code
                        if "一般名（全角）" in headers_ae:
                            new_row[headers_ae.index("一般名（全角）")] = name
                        for col_name in ae_columns:
                            if col_name in headers_ae:
                                new_row[headers_ae.index(col_name)] = (
                                    '○' if checked.get(col_name, False) else ''
                                )
                        if "出典" in headers_ae:
                            new_row[headers_ae.index("出典")] = source_text
                        if "登録日" in headers_ae:
                            new_row[headers_ae.index("登録日")] = today
                    else:
                        new_row = (
                            [code, name]
                            + ['○' if checked.get(c, False) else '' for c in ae_columns]
                            + [source_text, today]
                        )
                    ws_ae.append_row(new_row, value_input_option="USER_ENTERED")

                st.success(f"✅ {name} の副作用マスタを登録しました！")
                for col_name in ae_columns:
                    st.session_state.pop(f"maint_cb_{code}_{col_name}", None)
                st.session_state.pop(f"maint_src_{code}", None)
                st.session_state.pop("ae_pending_code", None)
                st.session_state.pop("ae_pending_name", None)
                st.session_state.pop("ae_select_target", None)
                fetch_sheet_realtime.clear()
                st.rerun()
            except Exception as e:
                st.error(f"❌ 登録エラー: {e}")

    with col_skip:
        if st.button(
            "⏭️ あとで登録する（スキップ）",
            use_container_width=True,
            key=f"maint_ae_skip_{code}",
        ):
            st.session_state.pop("ae_pending_code", None)
            st.session_state.pop("ae_pending_name", None)
            st.session_state.pop("ae_select_target", None)
            st.rerun()


def diagnose_pd_ae_alignment(pd_data, ae_data):
    if ae_data:
        all_keys = list(ae_data[0].keys())
    else:
        all_keys = []
    excluded = {"管理コード", "一般名（全角）", "出典", "登録日"}
    ae_columns = [k for k in all_keys if k not in excluded]

    result = {
        "symptom_matched": [], "symptom_unmatched": [],
        "drug_matched": [], "drug_unmatched": [], "no_type": [],
    }
    for row in pd_data:
        cat = str(row.get("カテゴリ名", "")).strip()
        cat_type = str(row.get("種別", "")).strip()
        if not cat:
            continue
        if cat_type == "症状群":
            (result["symptom_matched"] if cat in ae_columns
             else result["symptom_unmatched"]).append(cat)
        elif cat_type == "薬剤・薬効群":
            (result["drug_matched"] if cat in ae_columns
             else result["drug_unmatched"]).append(cat)
        else:
            result["no_type"].append(cat)
    return result, ae_columns


# ===== Streamlit UI =====
# ===== Streamlit UI =====
st.title("🧪 新薬メンテナンス")
st.caption("薬品マスタ・抗がん剤副作用マスタの新規登録・整備を行います")
st.divider()

if st.button("🔄 データを最新化する", key="btn_refresh_maint"):
    st.cache_data.clear()
    st.cache_resource.clear()
    st.rerun()

with st.spinner("データを読み込み中..."):
    drug_data, master_data, ae_data, pd_data = load_all_data()

ae_columns = get_ae_columns(ae_data)

# ===== ①薬品マスタ未登録チェック =====
st.subheader("① 薬品マスタ未登録チェック")
st.caption("薬剤情報シートに登場するが、薬品マスタに存在しない管理コードを警告します")

def normalize_for_match(text):
    table = str.maketrans('', '', ' 　')
    return str(text).translate(table).upper()

def search_master_candidates(product_name_raw, master_data):
    raw_norm = normalize_for_match(product_name_raw)
    hits = []
    for m in master_data:
        code = str(m.get("管理コード", "")).strip()
        if not code:
            continue
        for key in ("採用商品名（全角）", "一般名（全角）"):
            candidate = str(m.get(key, "")).strip()
            if candidate and len(candidate) >= 2 and normalize_for_match(candidate) in raw_norm:
                hits.append(m)
                break
    return hits

master_codes = {str(m.get('管理コード', '')).strip() for m in master_data}

def is_pending_code(code):
    """「要確認」の表記ゆれ（全角スペース混入等）を吸収して判定する"""
    normalized = re.sub(r'\s+', '', str(code))
    return normalized == '要確認'

# --- 「要確認」行（名称マッチング未済）の検出 ---
pending_match = {}
for d in drug_data:
    code = str(d.get('管理コード', '')).strip()
    if not is_pending_code(code):
        continue
    product_name = str(d.get('商品名', '')).strip()
    if not product_name:
        continue
    info = pending_match.setdefault(product_name, {
        'protocols': set(),
        'rows': [],
    })
    info['protocols'].add(str(d.get('プロトコールNo', '')).strip())
    info['rows'].append(d)

if pending_match:
    st.warning(f"⚠️ 管理コードが「要確認」の薬剤が {len(pending_match)} 種類あります")
    for product_name, info in pending_match.items():
        protocols_str = "、".join(sorted(info['protocols']))
        st.markdown(f"**{product_name}**（使用レジメン：{protocols_str}）")
        candidates = search_master_candidates(product_name, master_data)
        if candidates:
            for cand in candidates:
                cand_code = str(cand.get('管理コード', '')).strip()
                cand_name = str(cand.get('採用商品名（全角）', '') or cand.get('一般名（全角）', '')).strip()
                col_info, col_btn = st.columns([3, 1])
                with col_info:
                    st.caption(f"→ 薬品マスタに一致候補：**{cand_code}**（{cand_name}）")
                with col_btn:
                    if st.button(
                        "✅ 紐付ける",
                        key=f"btn_fix_pending_{product_name}_{cand_code}",
                    ):
                        try:
                            gc = get_gspread_client()
                            sh = gc.open_by_url(SPREADSHEET_URL)
                            ws_drug = sh.worksheet("薬剤情報")
                            all_vals = ws_drug.get_all_values()
                            headers_drug = all_vals[0]
                            code_col_idx = headers_drug.index('管理コード') + 1
                            name_col_idx = headers_drug.index('商品名') + 1
                            from openpyxl.utils import get_column_letter as gcl
                            updated_count = 0
                            for i, row in enumerate(all_vals[1:], start=2):
                                if (len(row) >= max(code_col_idx, name_col_idx)
                                        and row[code_col_idx - 1].strip() == '要確認'
                                        and row[name_col_idx - 1].strip() == product_name):
                                    ws_drug.update(
                                        range_name=f'{gcl(code_col_idx)}{i}',
                                        values=[[cand_code]],
                                    )
                                    updated_count += 1
                            st.success(
                                f"✅ 「{product_name}」の{updated_count}件を"
                                f"{cand_code}に紐付けました！"
                            )
                            fetch_sheet_realtime.clear()
                            st.rerun()
                        except Exception as e:
                            st.error(f"❌ 紐付けエラー: {e}")
        else:
            st.caption("　薬品マスタに自動一致候補が見つかりませんでした。")

            manual_mode_key = f"manual_search_mode_{product_name}"
            col_manual, col_new = st.columns(2)
            with col_manual:
                if st.button(
                    "🔍 既存の薬品マスタから手動で探す",
                    key=f"btn_manual_search_{product_name}",
                    use_container_width=True,
                ):
                    st.session_state[manual_mode_key] = True
                    st.rerun()
            with col_new:
                if st.button(
                    f"➕ 新規登録する",
                    key=f"btn_load_pending_{product_name}",
                    use_container_width=True,
                ):
                    st.session_state["newdrug_brand_prefill"] = product_name
                    st.session_state["newdrug_brand_full"] = product_name
                    st.session_state["scroll_hint_product"] = product_name
                    st.rerun()

            if st.session_state.get("scroll_hint_product") == product_name:
                st.success(
                    f"⬇️ 採用商品名「{product_name}」を"
                    "下の「② 新規薬剤登録」フォームに入力済みです。画面を下にスクロールしてください。"
                )
                if st.button(
                    "✅ この案内を閉じる",
                    key=f"btn_close_scroll_hint_{product_name}",
                ):
                    st.session_state.pop("scroll_hint_product", None)
                    st.rerun()

            if st.session_state.get(manual_mode_key):
                all_options = {
                    f"{str(m.get('管理コード','')).strip()}："
                    f"{str(m.get('一般名（全角）','')).strip()}"
                    f"（{str(m.get('採用商品名（全角）','')).strip()}）": m
                    for m in master_data
                    if str(m.get('管理コード', '')).strip()
                }
                manual_selected = st.selectbox(
                    f"「{product_name}」に対応する薬品マスタを検索・選択してください",
                    options=["選択してください"] + list(all_options.keys()),
                    key=f"manual_select_{product_name}",
                    placeholder="薬剤名を入力して検索...",
                )
                if manual_selected != "選択してください":
                    manual_cand = all_options[manual_selected]
                    manual_cand_code = str(manual_cand.get('管理コード', '')).strip()
                    if st.button(
                        f"✅ {manual_cand_code} に紐付ける",
                        key=f"btn_manual_fix_{product_name}_{manual_cand_code}",
                        type="primary",
                        use_container_width=True,
                    ):
                        try:
                            gc = get_gspread_client()
                            sh = gc.open_by_url(SPREADSHEET_URL)
                            ws_drug = sh.worksheet("薬剤情報")
                            all_vals = ws_drug.get_all_values()
                            headers_drug = all_vals[0]
                            code_col_idx = headers_drug.index('管理コード') + 1
                            name_col_idx = headers_drug.index('商品名') + 1
                            from openpyxl.utils import get_column_letter as gcl
                            updated_count = 0
                            for i, row in enumerate(all_vals[1:], start=2):
                                if (len(row) >= max(code_col_idx, name_col_idx)
                                        and is_pending_code(row[code_col_idx - 1])
                                        and row[name_col_idx - 1].strip() == product_name):
                                    ws_drug.update(
                                        range_name=f'{gcl(code_col_idx)}{i}',
                                        values=[[manual_cand_code]],
                                    )
                                    updated_count += 1
                            st.success(
                                f"✅ 「{product_name}」の{updated_count}件を"
                                f"{manual_cand_code}に紐付けました！"
                            )
                            st.session_state.pop(manual_mode_key, None)
                            fetch_sheet_realtime.clear()
                            st.rerun()
                        except Exception as e:
                            st.error(f"❌ 紐付けエラー: {e}")
else:
    st.success("✅ 「要確認」のままになっている薬剤はありません")

st.divider()

# --- 薬品マスタに存在しないコード（純粋な未登録）の検出 ---
missing = {}
for d in drug_data:
    code = str(d.get('管理コード', '')).strip()
    if not code or is_pending_code(code) or code in master_codes:
        continue
    info = missing.setdefault(code, {
        'name': str(d.get('商品名', '')).strip(),
        'protocols': set(),
    })
    info['protocols'].add(str(d.get('プロトコールNo', '')).strip())

# 診断用：missingに「要確認」系の文字列が紛れていないか確認表示
for _diag_code in list(missing.keys()):
    if '確認' in _diag_code:
        st.caption(
            f"🔍 診断：怪しいコードの文字コード一覧 → "
            f"{[hex(ord(c)) for c in _diag_code]}"
        )

if missing:
    st.warning(f"⚠️ 薬品マスタに未登録のコードが {len(missing)} 件あります")
    for code, info in missing.items():
        protocols_str = "、".join(sorted(info['protocols']))
        st.markdown(f"**{code}**：{info['name']}（使用レジメン：{protocols_str}）")
        if st.button(
            f"➕ {code} を登録フォームに読み込む",
            key=f"btn_load_missing_{code}",
        ):
            st.session_state["newdrug_fixed_code"] = code
            st.session_state["newdrug_brand_prefill"] = info['name']
            st.rerun()
else:
    st.success("✅ 薬品マスタ未登録のコードはありません")

st.divider()

# ===== ②新規薬剤登録フォーム =====
st.subheader("② 新規薬剤登録")

fixed_code = st.session_state.get("newdrug_fixed_code")

if fixed_code:
    st.info(f"🔧 薬剤情報シートに既存の管理コード「{fixed_code}」を登録します")
    new_code = fixed_code
    guessed_prefix = re.match(r'^[A-Z]+', fixed_code)
    prefix_for_kubun = guessed_prefix.group() if guessed_prefix else ""
    default_kubun = PREFIX_DEFAULT_KUBUN.get(prefix_for_kubun, "")
    if st.button("🔙 固定コード指定を解除する", key="btn_unfix_code"):
        st.session_state.pop("newdrug_fixed_code", None)
        st.session_state.pop("newdrug_brand_prefill", None)
        st.rerun()
brand_full = st.text_input(
    "① 採用商品名（全角）　※規格・銘柄名は省略し、一般名と同一表記で可",
    value=st.session_state.get("newdrug_brand_prefill", ""),
    key="newdrug_brand_full",
)

name_full_preview = st.text_input(
    "② 一般名（全角）　※入力すると既存マスタとの一致を自動検索します",
    key="newdrug_name_full",
)
if name_full_preview.strip():
    preview_candidates = search_master_candidates(name_full_preview, master_data)
    if preview_candidates:
        st.warning("⚠️ 薬品マスタに似た名称の薬剤が既に存在します。新規登録前にご確認ください：")
        for cand in preview_candidates:
            cand_code = str(cand.get('管理コード', '')).strip()
            cand_name = str(cand.get('採用商品名（全角）', '') or cand.get('一般名（全角）', '')).strip()
            st.caption(f"→ **{cand_code}**：{cand_name}（既存のこのコードを使用してください）")

if not fixed_code:
    prefix_label = st.selectbox(
        "③ 管理コードの分類を選択してください",
        options=list(PREFIX_OPTIONS.keys()),
        key="newdrug_prefix_label",
    )
    prefix = PREFIX_OPTIONS[prefix_label]
    prefix_for_kubun = prefix
    new_code = get_next_code(prefix, master_data)
    default_kubun = PREFIX_DEFAULT_KUBUN.get(prefix, "")
    st.success(f"📌 発番予定の管理コード：**{new_code}**　（薬剤区分：自動的に「{default_kubun}」になります）")

# 管理コードのプレフィックスに応じて薬効分類の選択肢を絞り込む
current_prefix = prefix_for_kubun if 'prefix_for_kubun' in dir() else ""
filtered_categories = sorted(set(
    str(m.get('薬効分類', '')).strip()
    for m in master_data
    if str(m.get('薬効分類', '')).strip()
    and re.match(rf'^{re.escape(current_prefix)}-?\d+$', str(m.get('管理コード', '')).strip())
)) if current_prefix else []

if not filtered_categories:
    filtered_categories = sorted(set(
        str(m.get('薬効分類', '')).strip()
        for m in master_data
        if str(m.get('薬効分類', '')).strip()
    ))

shiji_bunrui_options = get_unique_values(master_data, "支持療法分類")

with st.form(key="form_newdrug"):
    name_full = name_full_preview
    st.caption(
        f"採用商品名：**{brand_full or '（未入力）'}**　／　"
        f"一般名（全角）：**{name_full or '（未入力）'}**　／　管理コード：**{new_code}**"
    )

    category = st.selectbox(
        f"④ 薬効分類　※「{default_kubun or '該当分類'}」の既存リストから選択。"
        "無ければ下の「新しい薬効分類」に入力",
        options=["（選択してください）"] + filtered_categories,
        key="newdrug_category_select",
    )
    category_new = st.text_input(
        "新しい薬効分類（上で該当が無い場合のみ入力）",
        key="newdrug_category_new",
    )

    shiji_bunrui = st.selectbox(
        "支持療法分類（任意）",
        options=["（なし）"] + shiji_bunrui_options + ["その他"],
        key="newdrug_shiji_bunrui_select",
    )
    shiji_bunrui_other = ""
    if shiji_bunrui == "その他":
        shiji_bunrui_other = st.text_input(
            "支持療法分類（その他の内容を入力）",
            key="newdrug_shiji_bunrui_other",
        )

    with st.expander("詳細項目（任意・必要な場合のみ入力）"):
        tani = st.text_input("単位", key="newdrug_tani")
        toyokeiro = st.text_input("投与経路", key="newdrug_toyokeiro")
        kibo_eki = st.text_input("標準希釈液", key="newdrug_kibo_eki")
        filter_val = st.text_input("フィルター", key="newdrug_filter")
        shako = st.text_input("遮光", key="newdrug_shako")
        senpatsu = st.text_input("先発後発", key="newdrug_senpatsu")
        biko = st.text_input("備考", key="newdrug_biko")
        v_mg = st.text_input("1V当たりmg", key="newdrug_v_mg")
        kanja_setsumei = st.text_area("患者向け説明", key="newdrug_kanja_setsumei")
        seal_type = st.text_input("スケジュールシール用種類", key="newdrug_seal_type")
        alias = st.text_input("別名・旧採用品名", key="newdrug_alias")
        tanshuku = st.text_input("短縮注記", key="newdrug_tanshuku")

    if name_full:
        st.caption(f"💡 半角カナ変換プレビュー（参考・実際はASC関数で自動計算）：{preview_half_kana(name_full)}")

    submitted = st.form_submit_button(
        "✅ 薬品マスタに登録する",
        type="primary",
        use_container_width=True,
    )

if submitted:
    final_category = category_new.strip() if category_new.strip() else (
        category if category != "（選択してください）" else ""
    )
    final_kubun = default_kubun
    if shiji_bunrui == "その他":
        final_shiji_bunrui = shiji_bunrui_other.strip()
    elif shiji_bunrui == "（なし）":
        final_shiji_bunrui = ""
    else:
        final_shiji_bunrui = shiji_bunrui

    if not name_full.strip():
        st.error("⚠️ 一般名を入力してください")
    elif not final_category:
        st.error("⚠️ 薬効分類を選択または入力してください")
    else:
        try:
            gc = get_gspread_client()
            sh = gc.open_by_url(SPREADSHEET_URL)
            ws_master = sh.worksheet("薬品マスタ")

            row = [
                new_code,
                name_full.strip(),
                "",
                brand_full.strip() or name_full.strip(),
                "",
                final_category,
                final_kubun,
                final_shiji_bunrui,
                tani.strip() if 'tani' in dir() else "",
                toyokeiro.strip() if 'toyokeiro' in dir() else "",
                kibo_eki.strip() if 'kibo_eki' in dir() else "",
                filter_val.strip() if 'filter_val' in dir() else "",
                shako.strip() if 'shako' in dir() else "",
                senpatsu.strip() if 'senpatsu' in dir() else "",
                biko.strip() if 'biko' in dir() else "",
                v_mg.strip() if 'v_mg' in dir() else "",
                kanja_setsumei.strip() if 'kanja_setsumei' in dir() else "",
                seal_type.strip() if 'seal_type' in dir() else "",
                alias.strip() if 'alias' in dir() else "",
                tanshuku.strip() if 'tanshuku' in dir() else "",
            ]
            ws_master.append_row(row, value_input_option="USER_ENTERED")

            # 実際に追加された行番号を取得してから数式を書き込む（REF対策）
            all_codes = ws_master.col_values(1)
            actual_row = len(all_codes)  # append_row直後の最終行＝今追加した行
            ws_master.update(
                range_name=f'C{actual_row}',
                values=[[f"=ASC(B{actual_row})"]],
                value_input_option="USER_ENTERED",
            )
            ws_master.update(
                range_name=f'E{actual_row}',
                values=[[f"=ASC(D{actual_row})"]],
                value_input_option="USER_ENTERED",
            )

            st.success(f"✅ {new_code}（{name_full}）を薬品マスタに登録しました！")
            st.session_state.pop("newdrug_fixed_code", None)
            st.session_state.pop("newdrug_brand_prefill", None)

            if new_code.upper().startswith("AC"):
                st.session_state["ae_pending_code"] = new_code
                st.session_state["ae_pending_name"] = name_full.strip()

            fetch_sheet_realtime.clear()
            st.rerun()
        except Exception as e:
            st.error(f"❌ 登録エラー: {e}")

st.divider()

# ===== ③抗がん剤副作用マスタ整備 =====
st.subheader("③ 抗がん剤副作用マスタ整備")

if st.session_state.get("ae_pending_code"):
    show_ae_check_ui(
        st.session_state["ae_pending_code"],
        st.session_state["ae_pending_name"],
        ae_data, ae_columns,
    )
else:
    unregistered = get_unregistered_ae_drugs(master_data, ae_data)
    if unregistered:
        st.warning(f"⚠️ 副作用マスタ未登録（登録日が空）の抗がん剤が {len(unregistered)} 件あります")
        options = {f"{u['name']}（{u['code']}）": u for u in unregistered}
        selected_label = st.selectbox(
            "整備する薬剤を選択してください",
            options=["選択してください"] + list(options.keys()),
            key="ae_select_target",
        )
        if selected_label != "選択してください":
            target = options[selected_label]
            show_ae_check_ui(target['code'], target['name'], ae_data, ae_columns)
    else:
        st.success("✅ 抗がん剤副作用マスタ未登録の薬剤はありません")

st.divider()
st.divider()

# ===== おまけ：Pd整合性チェック =====
st.subheader("🔧 おまけ：Pd整合性チェック")
st.caption("Pdシートの「種別」列に基づき、抗がん剤副作用マスタとの整合性を確認します（独立したメンテナンス機能です）")

if st.button("🔍 整合性をチェックする", key="btn_check_pd_alignment_maint"):
    diag_result, _ = diagnose_pd_ae_alignment(pd_data, ae_data)
    st.session_state["pd_diagnosis_maint"] = diag_result
    st.rerun()

if "pd_diagnosis_maint" in st.session_state:
    diag = st.session_state["pd_diagnosis_maint"]

    st.markdown("**✅ 症状群カテゴリ（マスタと一致）**")
    st.write(diag["symptom_matched"] if diag["symptom_matched"] else "（なし）")

    st.markdown("**✅ 薬剤・薬効群カテゴリ（マスタと一致）**")
    st.write(diag["drug_matched"] if diag["drug_matched"] else "（なし）")

    st.markdown("**⚠️ 未対応：症状群カテゴリ（手動登録が必要）**")
    if diag["symptom_unmatched"]:
        for cat in diag["symptom_unmatched"]:
            st.warning(
                f"「{cat}」列が抗がん剤副作用マスタにありません。"
                f"[スプレッドシートを開く]({SPREADSHEET_URL})で"
                f"「{cat}」列を末尾に追加し、各薬剤を確認して該当するものに○をつけ、"
                f"出典・登録日を記入してください。"
            )
    else:
        st.success("未対応の症状群カテゴリはありません")

    st.markdown("**⚠️ 未対応：薬剤・薬効群カテゴリ（手動登録が必要）**")
    if diag["drug_unmatched"]:
        for cat in diag["drug_unmatched"]:
            st.warning(
                f"「{cat}」列が抗がん剤副作用マスタにありません。"
                f"[スプレッドシートを開く]({SPREADSHEET_URL})で"
                f"「{cat}」列を末尾に追加し、各薬剤を確認して該当するものに○をつけ、"
                f"出典・登録日を記入してください。"
            )
    else:
        st.success("未対応の薬剤・薬効群カテゴリはありません")

    if diag["no_type"]:
        st.markdown("**❓ 種別未設定のカテゴリ**")
        st.write(diag["no_type"])
        st.caption("Pdシートの「種別」列に「症状群」または「薬剤・薬効群」を設定してください。")
